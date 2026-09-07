"""F21 acquisition gate: validate registered bytes and produce a machine-readable record.

The validator MUST hash the local file first and abort before opening the workbook
if the hash does not match the preregistered SHA-256. A mismatch is a source-version
discrepancy, not a reason to substitute the file.

Produces data/acquisition/WUP2025-F21-validation.json.

Two commit SHAs are recorded:
  validator_commit  -- the commit of this validator at the time of the run.
                       Advances with each acquisition-gate PR revision.
  construction_commit -- the builder commit (70249dad) that will be used for
                         panel construction. Fixed independently of validator
                         revisions so source validation and panel construction
                         remain separately attributable.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import UTC, datetime
from enum import Enum
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, ConfigDict

REGISTERED_SHA256 = "3a96030d87aec6c1c50f658d5321067d6345e1ab936c5d2854524f972caa75c0"
REGISTERED_BYTES = 10_671_837
REGISTERED_SOURCE_ID = "un_wup_2025_cities"
REGISTERED_TABLE = "F21"
REGISTERED_VINTAGE = "WUP_2025_F21"
REQUIRED_COLUMNS = frozenset(
    {
        "LocID",
        "ISO3_Code",
        "City_Code",
        "City_Name",
        "PWCent_Longitude",
        "PWCent_Latitude",
        "1975",
        "2025",
        "2050",
    }
)
REQUIRED_SHEET = "Data"

# Builder lineage fixed at the post-#212 merge commit.
CONSTRUCTION_COMMIT = "70249dad153f4ba864fd7d566d05893be2421813"


class EmpiricalActivationState(str, Enum):
    blocked_source_unavailable = "blocked_source_unavailable"
    blocked_source_identity = "blocked_source_identity"
    blocked_source_validation = "blocked_source_validation"
    ready_for_panel_construction = "ready_for_panel_construction"
    panel_frozen = "panel_frozen"
    ready_for_performance_evaluation = "ready_for_performance_evaluation"
    evaluation_complete = "evaluation_complete"


class LicensingDisposition(BaseModel):
    model_config = ConfigDict(extra="forbid")
    license: str
    redistribution_permitted: bool
    redistribution_notes: str
    repository_policy: str
    repository_policy_notes: str


class IdentityCheck(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expected_sha256: str
    observed_sha256: str
    expected_bytes: int
    observed_bytes: int
    passed: bool


class SchemaCheck(BaseModel):
    model_config = ConfigDict(extra="forbid")
    required_sheet_present: bool
    required_sheet: str
    required_columns_present: bool
    required_columns: list[str]
    missing_columns: list[str]
    passed: bool


class ContentCheck(BaseModel):
    model_config = ConfigDict(extra="forbid")
    year_columns_observed: list[int]
    estimate_cutoff_year: int
    total_rows: int
    city_count: int
    country_count: int
    duplicate_city_id_count: int
    null_population_rows: int
    below_threshold_populated_rows: int
    estimate_rows: int
    projection_rows: int
    passed: bool


class PanelFencePRA(BaseModel):
    """Panel Freeze PR specification: diagnostics only, no performance output."""

    model_config = ConfigDict(extra="forbid")
    permitted: list[str]
    prohibited: list[str]
    acceptance: str


class PanelFencePRB(BaseModel):
    """Performance/W2 PR specification: gated on PR A merge."""

    model_config = ConfigDict(extra="forbid")
    prerequisite: str
    contains: list[str]


class PanelFenceProtocol(BaseModel):
    """Machine-readable two-stage PR fence. CI enforces permitted/prohibited lists."""

    model_config = ConfigDict(extra="forbid")
    pr_a_panel_freeze: PanelFencePRA
    pr_b_performance: PanelFencePRB


_FENCE = PanelFenceProtocol(
    pr_a_panel_freeze=PanelFencePRA(
        permitted=[
            "origin_eligibility_table",
            "singleton_exclusion_counts",
            "city_and_country_counts_by_origin",
            "paired_row_counts",
            "weight_sums_and_distribution",
            "leakage_assertions",
            "schema_checks",
            "panel_sha256",
        ],
        prohibited=[
            "b0_errors",
            "b1_errors",
            "rmse",
            "mae",
            "relative_rmse_improvement",
            "mae_difference",
            "gate_evaluation",
            "model_preference_language",
        ],
        acceptance="recorded_checklist_decision_required_before_merge",
    ),
    pr_b_performance=PanelFencePRB(
        prerequisite="pr_a_merged",
        contains=[
            "performance_metrics",
            "w2_inconclusive_probe_registration",
            "w2_complete_result_registration",
            "gate_resolution",
        ],
    ),
)

_LICENSING = LicensingDisposition(
    license="CC BY 3.0 IGO",
    redistribution_permitted=True,
    redistribution_notes=(
        "CC BY 3.0 IGO permits redistribution with attribution. "
        "UN copyright notice and citation required."
    ),
    repository_policy="review_before_commit",
    repository_policy_notes=(
        "Raw workbook intentionally not committed pending project "
        "review-before-commit policy. Legal permission and repository "
        "policy are distinct: redistribution is permitted under CC BY 3.0 IGO; "
        "committing to this repository requires explicit review per CONTRIBUTING.md."
    ),
)


class F21AcquisitionValidation(BaseModel):
    model_config = ConfigDict(extra="forbid")
    source_id: str
    table: str
    vintage: str
    local_path: str
    adapter_module: str
    # validator_commit: commit of this validator module at run time.
    # construction_commit: fixed builder lineage (post-#212); invariant across
    # validator revisions so validation and construction remain separately attributable.
    validator_commit: str
    construction_commit: str
    validated_at: str
    identity: IdentityCheck
    licensing: LicensingDisposition
    schema_check: SchemaCheck
    content: ContentCheck
    overall_outcome: Literal["pass", "fail"]
    activation_state: EmpiricalActivationState
    panel_fence_protocol: PanelFenceProtocol


def _sha256_and_size(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    total = 0
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
            total += len(chunk)
    return digest.hexdigest(), total


def _empty_content() -> ContentCheck:
    return ContentCheck(
        year_columns_observed=[],
        estimate_cutoff_year=2025,
        total_rows=0,
        city_count=0,
        country_count=0,
        duplicate_city_id_count=0,
        null_population_rows=0,
        below_threshold_populated_rows=0,
        estimate_rows=0,
        projection_rows=0,
        passed=False,
    )


def validate_f21(
    local_path: Path,
    validator_commit: str,
) -> F21AcquisitionValidation:
    """Validate the registered F21 bytes and produce a structured record.

    Aborts before opening the workbook if the SHA-256 does not match.
    validator_commit identifies this validator run; construction_commit is
    fixed to the post-#212 builder lineage and does not change with validator
    revisions.
    """
    observed_sha, observed_bytes = _sha256_and_size(local_path)
    identity_passed = (
        observed_sha == REGISTERED_SHA256 and observed_bytes == REGISTERED_BYTES
    )
    identity = IdentityCheck(
        expected_sha256=REGISTERED_SHA256,
        observed_sha256=observed_sha,
        expected_bytes=REGISTERED_BYTES,
        observed_bytes=observed_bytes,
        passed=identity_passed,
    )

    common = {
        "source_id": REGISTERED_SOURCE_ID,
        "table": REGISTERED_TABLE,
        "vintage": REGISTERED_VINTAGE,
        "local_path": str(local_path),
        "adapter_module": "urban_growth.f21_acquisition_gate",
        "validator_commit": validator_commit,
        "construction_commit": CONSTRUCTION_COMMIT,
        "validated_at": datetime.now(UTC).isoformat(),
        "identity": identity,
        "licensing": _LICENSING,
        "panel_fence_protocol": _FENCE,
    }

    if not identity_passed:
        return F21AcquisitionValidation(
            **common,
            schema_check=SchemaCheck(
                required_sheet_present=False,
                required_sheet=REQUIRED_SHEET,
                required_columns_present=False,
                required_columns=sorted(REQUIRED_COLUMNS),
                missing_columns=[],
                passed=False,
            ),
            content=_empty_content(),
            overall_outcome="fail",
            activation_state=EmpiricalActivationState.blocked_source_identity,
        )

    # Identity confirmed — open workbook for schema check.
    import openpyxl

    wb = openpyxl.load_workbook(str(local_path), read_only=True, data_only=True)
    sheet_present = REQUIRED_SHEET in wb.sheetnames
    missing_cols: list[str] = []
    schema_passed = False
    if sheet_present:
        ws = wb[REQUIRED_SHEET]
        header = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in next(ws.iter_rows(max_row=1))
        ]
        missing_cols = sorted(REQUIRED_COLUMNS - set(header))
        schema_passed = not missing_cols
    wb.close()

    schema_check = SchemaCheck(
        required_sheet_present=sheet_present,
        required_sheet=REQUIRED_SHEET,
        required_columns_present=schema_passed,
        required_columns=sorted(REQUIRED_COLUMNS),
        missing_columns=missing_cols,
        passed=sheet_present and schema_passed,
    )

    if not schema_check.passed:
        return F21AcquisitionValidation(
            **common,
            schema_check=schema_check,
            content=_empty_content(),
            overall_outcome="fail",
            activation_state=EmpiricalActivationState.blocked_source_validation,
        )

    # Schema confirmed — content checks via production adapter.
    from urban_growth.adapters.wup import read_f21_city_population

    panel = read_f21_city_population(str(local_path))

    wb2 = openpyxl.load_workbook(str(local_path), read_only=True, data_only=True)
    ws2 = wb2[REQUIRED_SHEET]
    raw_header = [
        str(c.value).strip() if c.value is not None else ""
        for c in next(ws2.iter_rows(max_row=1))
    ]
    year_cols = sorted(int(h) for h in raw_header if re.fullmatch(r"\d{4}", h))
    wb2.close()

    dup_city_ids = int(
        panel.groupby("city_id")["year"].count().gt(panel["year"].nunique()).sum()
    )
    content = ContentCheck(
        year_columns_observed=year_cols,
        estimate_cutoff_year=2025,
        total_rows=len(panel),
        city_count=int(panel["city_id"].nunique()),
        country_count=int(panel["ISO3_Code"].nunique()),
        duplicate_city_id_count=dup_city_ids,
        null_population_rows=0,  # adapter rejects nulls; reaching here means 0
        below_threshold_populated_rows=0,  # adapter rejects below-threshold
        estimate_rows=int((panel["observation_type"] == "estimate").sum()),
        projection_rows=int((panel["observation_type"] == "projection").sum()),
        passed=True,
    )

    return F21AcquisitionValidation(
        **common,
        schema_check=schema_check,
        content=content,
        overall_outcome="pass",
        activation_state=EmpiricalActivationState.ready_for_panel_construction,
    )


def _parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Validate registered F21 bytes and write acquisition gate record"
    )
    p.add_argument("--f21", required=True, type=Path, help="Local path to F21 workbook")
    p.add_argument(
        "--validator-commit",
        required=True,
        help="Git commit SHA of this validator at run time",
    )
    p.add_argument(
        "--output",
        type=Path,
        default=Path("data/acquisition/WUP2025-F21-validation.json"),
        help="Output path for validation JSON",
    )
    return p


def main() -> None:
    args = _parser().parse_args()
    record = validate_f21(args.f21, args.validator_commit)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(record.model_dump(), indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(record.model_dump(), indent=2, sort_keys=True))
    if record.overall_outcome != "pass":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
